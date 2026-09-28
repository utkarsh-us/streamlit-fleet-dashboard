import streamlit as st
import pandas as pd
import os
import re
import io
import time
import datetime
import warnings
import openpyxl
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# ==================================================
# PAGE CONFIG
# ==================================================
st.set_page_config(
    page_title="Fleet Compliance Dashboard",
    page_icon="🚛",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ==================================================
# CSS — ULTRA MODERN UI
# ==================================================
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&display=swap');

html, body, [class*="css"] {
    font-family: 'Inter', 'Segoe UI', -apple-system, sans-serif;
    -webkit-font-smoothing: antialiased;
    -moz-osx-font-smoothing: grayscale;
}

/* ============ ANIMATED PAGE FADE-IN ============ */
@keyframes fadeInUp {
    from { opacity: 0; transform: translateY(12px); }
    to   { opacity: 1; transform: translateY(0); }
}
@keyframes fadeIn {
    from { opacity: 0; }
    to   { opacity: 1; }
}
@keyframes slideInLeft {
    from { opacity: 0; transform: translateX(-8px); }
    to   { opacity: 1; transform: translateX(0); }
}
@keyframes pulse {
    0%, 100% { opacity: 1; }
    50%      { opacity: 0.55; }
}
@keyframes shimmer {
    0%   { background-position: -1000px 0; }
    100% { background-position: 1000px 0; }
}

/* Fade in main content */
.main .block-container > div {
    animation: fadeInUp 0.5s cubic-bezier(0.4, 0, 0.2, 1);
}
section[data-testid="stSidebar"] > div {
    animation: slideInLeft 0.4s cubic-bezier(0.4, 0, 0.2, 1);
}

/* ============ HIDE STREAMLIT CHROME ============ */
#MainMenu {visibility: hidden;}
footer {visibility: hidden;}
header[data-testid="stHeader"] { background: transparent !important; box-shadow: none !important; }
button[data-testid="baseButton-headerNoPadding"] {
    visibility: visible !important;
    opacity: 1 !important;
    color: #0f172a !important;
}

/* ============ SOFT BACKGROUND ============ */
.stApp {
    background:
        radial-gradient(circle at 15% 10%, rgba(14,165,233,0.08), transparent 40%),
        radial-gradient(circle at 85% 5%, rgba(99,102,241,0.06), transparent 40%),
        linear-gradient(180deg, #f8fafc 0%, #eef2f7 100%);
}

.block-container {
    padding-top: 2rem !important;
    padding-bottom: 0.5rem !important;
    max-width: 100% !important;
}
section.main > div { padding-bottom: 0 !important; }
div[data-testid="stAppViewContainer"] > .main { padding-bottom: 0 !important; }
.element-container:has(div[data-testid="stDataFrame"]) { margin-bottom: 0 !important; }
div[data-testid="stDataFrame"] { margin-bottom: 0 !important; }

/* ============ KPI CARDS — GLASSY + ANIMATED ============ */
[data-testid="stMetric"] {
    background: linear-gradient(135deg, rgba(255,255,255,0.9) 0%, rgba(248,250,252,0.85) 100%);
    backdrop-filter: blur(12px);
    border: 1px solid rgba(226,232,240,0.9);
    border-radius: 16px;
    padding: 20px 24px;
    box-shadow:
        0 1px 2px rgba(15,23,42,0.04),
        0 8px 24px -8px rgba(15,23,42,0.08);
    transition: transform 0.25s cubic-bezier(0.4, 0, 0.2, 1),
                box-shadow 0.25s cubic-bezier(0.4, 0, 0.2, 1),
                border-color 0.25s ease;
    position: relative;
    overflow: hidden;
    animation: fadeInUp 0.6s cubic-bezier(0.4, 0, 0.2, 1) backwards;
}
[data-testid="stMetric"]::before {
    content: '';
    position: absolute;
    top: 0; left: 0; right: 0;
    height: 3px;
    background: linear-gradient(90deg, #0ea5e9, #6366f1, #8b5cf6);
    opacity: 0.85;
}
[data-testid="stMetric"]:hover {
    transform: translateY(-4px);
    box-shadow:
        0 4px 8px rgba(15,23,42,0.06),
        0 20px 40px -12px rgba(14,165,233,0.25);
    border-color: #bae6fd;
}
[data-testid="stMetricLabel"] {
    font-size: 12px !important;
    font-weight: 700 !important;
    color: #64748b !important;
    text-transform: uppercase;
    letter-spacing: 1px;
    display: flex;
    align-items: center;
    gap: 6px;
}
[data-testid="stMetricValue"] {
    font-size: 32px !important;
    font-weight: 800 !important;
    color: #0f172a !important;
    letter-spacing: -1px;
}

/* Staggered animation delays */
[data-testid="stMetric"]:nth-of-type(1) { animation-delay: 0.05s; }
[data-testid="stMetric"]:nth-of-type(2) { animation-delay: 0.10s; }
[data-testid="stMetric"]:nth-of-type(3) { animation-delay: 0.15s; }
[data-testid="stMetric"]:nth-of-type(4) { animation-delay: 0.20s; }

/* ============ SKELETON LOADER ============ */
.skeleton {
    background: linear-gradient(
        90deg,
        #f1f5f9 0%,
        #e2e8f0 50%,
        #f1f5f9 100%
    );
    background-size: 1000px 100%;
    animation: shimmer 1.6s linear infinite;
    border-radius: 12px;
}
.skeleton-card {
    height: 100px;
    margin-bottom: 16px;
}
.skeleton-row {
    height: 40px;
    margin-bottom: 8px;
}
.skeleton-title {
    height: 32px;
    width: 40%;
    margin-bottom: 20px;
}

/* ============ LUCIDE-STYLE ICON ============ */
.icon {
    display: inline-flex;
    align-items: center;
    justify-content: center;
    width: 18px;
    height: 18px;
    flex-shrink: 0;
    vertical-align: middle;
}

.icon svg {
    width: 100%;
    height: 100%;
    stroke: currentColor;
    stroke-width: 2;
    stroke-linecap: round;
    stroke-linejoin: round;
    fill: none;
}

.icon-blue    { color: #0ea5e9; }
.icon-green   { color: #10b981; }
.icon-yellow  { color: #f59e0b; }
.icon-red     { color: #ef4444; }
.icon-slate   { color: #64748b; }
.icon-violet  { color: #8b5cf6; }

/* ============ SECTION HEADERS WITH ICONS ============ */
.section-header {
    display: flex;
    align-items: center;
    gap: 10px;
    font-size: 20px;
    font-weight: 800;
    color: #0f172a;
    letter-spacing: -0.3px;
    margin: 8px 0 16px 0;
    padding-bottom: 8px;
    border-bottom: 2px solid #f1f5f9;
}
.section-header .icon {
    width: 24px;
    height: 24px;
}

/* ============ DIALOG ============ */
div[data-testid="stDialog"] > div:first-child {
    border-radius: 20px !important;
    box-shadow: 0 40px 80px -20px rgba(15,23,42,0.35) !important;
    border: 1px solid rgba(226,232,240,0.9) !important;
    background: #ffffff !important;
    padding: 12px 8px !important;
    backdrop-filter: blur(20px);
    animation: fadeInUp 0.35s cubic-bezier(0.4, 0, 0.2, 1);
}
div[data-testid="stDialog"] h2 {
    font-size: 22px !important;
    font-weight: 800 !important;
    color: #0f172a !important;
    letter-spacing: -0.5px;
    padding-bottom: 14px;
    border-bottom: 1px solid #f1f5f9;
    margin-bottom: 18px !important;
}
div[data-testid="stDialog"] h3 {
    font-size: 13px !important;
    font-weight: 800 !important;
    color: #334155 !important;
    text-transform: uppercase;
    letter-spacing: 1px;
    margin-bottom: 16px !important;
    padding-bottom: 8px;
    border-bottom: 2px solid #0ea5e9;
    display: inline-block;
}
div[data-testid="stDialog"] input,
div[data-testid="stDialog"] textarea {
    background-color: #f8fafc !important;
    border: 1px solid #e2e8f0 !important;
    border-radius: 10px !important;
    color: #0f172a !important;
    font-size: 14px !important;
    font-weight: 500 !important;
    padding: 11px 14px !important;
    -webkit-text-fill-color: #0f172a !important;
    opacity: 1 !important;
}
div[data-testid="stDialog"] label {
    font-size: 10px !important;
    font-weight: 700 !important;
    color: #64748b !important;
    text-transform: uppercase;
    letter-spacing: 0.8px;
    margin-bottom: 4px !important;
}
div[data-testid="stDialog"] input:disabled {
    color: #0f172a !important;
    -webkit-text-fill-color: #0f172a !important;
    background-color: #f8fafc !important;
    opacity: 1 !important;
    cursor: default !important;
}
div[data-testid="stDialog"] .stButton > button,
div[data-testid="stDialog"] .stLinkButton > a {
    border-radius: 10px !important;
    font-weight: 600 !important;
    font-size: 13px !important;
    padding: 11px 18px !important;
    border: 1px solid #e2e8f0 !important;
    background: #ffffff !important;
    color: #0f172a !important;
    transition: all 0.2s cubic-bezier(0.4, 0, 0.2, 1) !important;
    text-align: center !important;
    text-decoration: none !important;
}
div[data-testid="stDialog"] .stButton > button:hover,
div[data-testid="stDialog"] .stLinkButton > a:hover {
    background: linear-gradient(135deg, #0ea5e9 0%, #0284c7 100%) !important;
    color: #ffffff !important;
    border-color: transparent !important;
    transform: translateY(-2px);
    box-shadow: 0 8px 20px -4px rgba(14,165,233,0.45);
}

/* ============ DATAFRAME ============ */
div[data-testid="stDataFrame"] {
    border-radius: 14px;
    overflow: hidden;
    border: 1px solid rgba(226,232,240,0.9);
    box-shadow: 0 1px 3px rgba(15,23,42,0.04), 0 8px 24px -8px rgba(15,23,42,0.06);
    background: #ffffff;
}
div[data-testid="stDataFrame"] td:first-child {
    color: #0ea5e9 !important;
    font-weight: 600 !important;
}
div[data-testid="stDataFrame"] input[type="checkbox"] {
    accent-color: #0ea5e9 !important;
}

/* ============ SIDEBAR ============ */
section[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #ffffff 0%, #f8fafc 60%, #eef2f7 100%);
    border-right: 1px solid #e2e8f0;
    box-shadow: 4px 0 24px -12px rgba(15,23,42,0.08);
}
section[data-testid="stSidebar"] h3 {
    font-size: 20px !important;
    font-weight: 900 !important;
    color: #0f172a !important;
    letter-spacing: -0.5px;
    margin-top: 8px !important;
    margin-bottom: 22px !important;
    padding-bottom: 14px;
    border-bottom: 3px solid transparent;
    background-image: linear-gradient(90deg, #0ea5e9, #6366f1);
    background-size: 60px 3px;
    background-repeat: no-repeat;
    background-position: 0 100%;
    display: inline-block;
}
section[data-testid="stSidebar"] label {
    font-size: 10px !important;
    font-weight: 800 !important;
    color: #64748b !important;
    text-transform: uppercase;
    letter-spacing: 1px;
    margin-bottom: 8px !important;
}

/* ============ RADIO PILLS ============ */
section[data-testid="stSidebar"] div[role="radiogroup"] {
    display: flex;
    gap: 8px;
    margin-top: 4px;
}
section[data-testid="stSidebar"] div[role="radiogroup"] > label {
    flex: 1;
    background: #ffffff;
    border: 1.5px solid #e2e8f0;
    border-radius: 12px;
    padding: 11px 10px !important;
    text-align: center;
    cursor: pointer;
    transition: all 0.2s cubic-bezier(0.4, 0, 0.2, 1);
    font-size: 12px !important;
    font-weight: 800 !important;
    color: #334155 !important;
    text-transform: uppercase !important;
    letter-spacing: 0.6px !important;
    margin: 0 !important;
    white-space: nowrap !important;
    overflow: hidden !important;
    box-shadow: 0 1px 2px rgba(15,23,42,0.03);
}
section[data-testid="stSidebar"] div[role="radiogroup"] > label:hover {
    border-color: #0ea5e9;
    background: #f0f9ff;
    color: #0284c7 !important;
    transform: translateY(-1px);
}
section[data-testid="stSidebar"] div[role="radiogroup"] > label:has(input:checked),
section[data-testid="stSidebar"] div[role="radiogroup"] > label[data-checked="true"] {
    background: linear-gradient(135deg, #0ea5e9 0%, #0284c7 100%) !important;
    border-color: transparent !important;
    color: #ffffff !important;
    box-shadow: 0 6px 16px -4px rgba(14,165,233,0.5);
}
section[data-testid="stSidebar"] div[role="radiogroup"] > label > div:first-child,
section[data-testid="stSidebar"] div[role="radiogroup"] > label > div:first-child * {
    display: none !important;
    width: 0 !important;
    height: 0 !important;
    margin: 0 !important;
}
section[data-testid="stSidebar"] input[type="radio"] { display: none !important; }

/* ============ SELECTBOX / INPUT ============ */
section[data-testid="stSidebar"] div[data-baseweb="select"] > div {
    background: #ffffff !important;
    border: 1.5px solid #e2e8f0 !important;
    border-radius: 12px !important;
    transition: all 0.2s ease;
    font-weight: 500;
    box-shadow: 0 1px 2px rgba(15,23,42,0.03);
}
section[data-testid="stSidebar"] div[data-baseweb="select"] > div:hover {
    border-color: #0ea5e9 !important;
    box-shadow: 0 0 0 4px rgba(14,165,233,0.12);
}
section[data-testid="stSidebar"] input[type="text"] {
    background: #ffffff !important;
    border: 1.5px solid #e2e8f0 !important;
    border-radius: 12px !important;
    color: #0f172a !important;
    font-size: 14px !important;
    padding: 11px 14px !important;
    -webkit-text-fill-color: #0f172a !important;
    box-shadow: 0 1px 2px rgba(15,23,42,0.03);
    transition: all 0.2s ease;
}
section[data-testid="stSidebar"] input[type="text"]:focus {
    border-color: #0ea5e9 !important;
    box-shadow: 0 0 0 4px rgba(14,165,233,0.15) !important;
}

/* ============ RESET BUTTON ============ */
section[data-testid="stSidebar"] .stButton > button {
    background: linear-gradient(135deg, #0ea5e9 0%, #0284c7 100%) !important;
    border: none !important;
    color: #ffffff !important;
    border-radius: 12px !important;
    font-weight: 800 !important;
    font-size: 13px !important;
    padding: 13px 18px !important;
    letter-spacing: 0.5px;
    transition: all 0.25s cubic-bezier(0.4, 0, 0.2, 1) !important;
    box-shadow: 0 6px 16px -4px rgba(14,165,233,0.4);
}
section[data-testid="stSidebar"] .stButton > button:hover {
    transform: translateY(-2px);
    box-shadow: 0 10px 24px -4px rgba(14,165,233,0.5) !important;
    background: linear-gradient(135deg, #0284c7 0%, #0369a1 100%) !important;
}
section[data-testid="stSidebar"] hr {
    border-color: #e2e8f0 !important;
    margin: 22px 0 !important;
}

/* ============ LAST UPDATED ============ */
.small-updated {
    text-align: left;
    font-size: 11px;
    color: #94a3b8;
    padding: 14px 0 0 0;
    letter-spacing: 0.4px;
    line-height: 1.6;
}
.small-updated strong {
    color: #475569;
    font-weight: 700;
}

/* ============ HEADINGS ============ */
h1 {
    font-size: 28px !important;
    font-weight: 900 !important;
    color: #0f172a !important;
    letter-spacing: -0.6px;
}
h3 {
    font-weight: 800 !important;
    letter-spacing: -0.3px;
    color: #0f172a !important;
}

/* ============ VIEW BUTTON IN TABLE ============ */
div[data-testid="stButton"] > button.view-btn {
    background: linear-gradient(135deg, #0ea5e9 0%, #0284c7 100%) !important;
    color: #ffffff !important;
    border: none !important;
    border-radius: 8px !important;
    padding: 5px 16px !important;
    font-size: 12px !important;
    font-weight: 700 !important;
    min-height: 30px !important;
    height: 30px !important;
    box-shadow: 0 3px 8px -2px rgba(14,165,233,0.4) !important;
    transition: all 0.2s cubic-bezier(0.4, 0, 0.2, 1) !important;
}
div[data-testid="stButton"] > button.view-btn:hover {
    transform: translateY(-1px) !important;
    box-shadow: 0 6px 16px -2px rgba(14,165,233,0.55) !important;
    background: linear-gradient(135deg, #0284c7 0%, #0369a1 100%) !important;
}

/* ============ PRIMARY BUTTONS ============ */
button[kind="primary"],
.stButton > button[kind="primary"] {
    background: linear-gradient(135deg, #0ea5e9 0%, #0284c7 100%) !important;
    border-color: transparent !important;
    color: #ffffff !important;
    box-shadow: 0 6px 16px -4px rgba(14,165,233,0.4) !important;
}

/* ============ SCROLLBAR ============ */
::-webkit-scrollbar {
    width: 10px;
    height: 10px;
}
::-webkit-scrollbar-track {
    background: transparent;
}
::-webkit-scrollbar-thumb {
    background: linear-gradient(180deg, #cbd5e1, #94a3b8);
    border-radius: 10px;
    border: 2px solid #f8fafc;
}
::-webkit-scrollbar-thumb:hover {
    background: linear-gradient(180deg, #94a3b8, #64748b);
}
</style>
""", unsafe_allow_html=True)


# ==================================================
# ICON HELPERS (LUCIDE-STYLE SVGs)
# ==================================================
def icon(name, color="slate", size=18):
    """
    Returns HTML for a Lucide-style icon.
    Colors: blue, green, yellow, red, slate, violet
    """
    paths = {
        "search":     '<circle cx="11" cy="11" r="8"/><path d="m21 21-4.3-4.3"/>',
        "filter":     '<polygon points="22 3 2 3 10 12.46 10 19 14 21 14 12.46 22 3"/>',
        "calendar":   '<rect width="18" height="18" x="3" y="4" rx="2"/><path d="M16 2v4M8 2v4M3 10h18"/>',
        "truck":      '<path d="M14 18V6a2 2 0 0 0-2-2H4a2 2 0 0 0-2 2v11a1 1 0 0 0 1 1h2"/><path d="M15 18H9"/><path d="M19 18h2a1 1 0 0 0 1-1v-3.65a1 1 0 0 0-.22-.624l-3.48-4.35A1 1 0 0 0 17.52 8H14"/><circle cx="17" cy="18" r="2"/><circle cx="7" cy="18" r="2"/>',
        "check-circle": '<path d="M22 11.08V12a10 10 0 1 1-5.93-9.14"/><path d="m9 11 3 3L22 4"/>',
        "alert":      '<path d="m21.73 18-8-14a2 2 0 0 0-3.48 0l-8 14A2 2 0 0 0 4 21h16a2 2 0 0 0 1.73-3Z"/><path d="M12 9v4M12 17h.01"/>',
        "x-circle":   '<circle cx="12" cy="12" r="10"/><path d="m15 9-6 6M9 9l6 6"/>',
        "file-text":  '<path d="M15 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V7Z"/><path d="M14 2v4a2 2 0 0 0 2 2h4"/><path d="M10 9H8M16 13H8M16 17H8"/>',
        "file":       '<path d="M15 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V7Z"/><path d="M14 2v4a2 2 0 0 0 2 2h4"/>',
        "eye":        '<path d="M2 12s3-7 10-7 10 7 10 7-3 7-10 7-10-7-10-7Z"/><circle cx="12" cy="12" r="3"/>',
        "chart":      '<path d="M3 3v18h18"/><path d="M18 17V9M13 17V5M8 17v-3"/>',
        "refresh":    '<path d="M3 12a9 9 0 0 1 9-9 9.75 9.75 0 0 1 6.74 2.74L21 8"/><path d="M21 3v5h-5"/><path d="M21 12a9 9 0 0 1-9 9 9.75 9.75 0 0 1-6.74-2.74L3 16"/><path d="M8 16H3v5"/>',
        "clock":      '<circle cx="12" cy="12" r="10"/><path d="M12 6v6l4 2"/>',
        "shield":     '<path d="M20 13c0 5-3.5 7.5-7.66 8.95a1 1 0 0 1-.67-.01C7.5 20.5 4 18 4 13V6a1 1 0 0 1 1-1c2 0 4.5-1.2 6.24-2.72a1.17 1.17 0 0 1 1.52 0C14.51 3.81 17 5 19 5a1 1 0 0 1 1 1z"/>',
        "car":        '<path d="M19 17h2c.6 0 1-.4 1-1v-3c0-.9-.7-1.7-1.5-1.9C18.7 10.6 16 10 16 10s-1.3-1.4-2.2-2.3c-.5-.4-1.1-.7-1.8-.7H5c-.6 0-1.1.4-1.4.9l-1.4 2.9A3.7 3.7 0 0 0 2 12v4c0 .6.4 1 1 1h2"/><circle cx="7" cy="17" r="2"/><path d="M9 17h6"/><circle cx="17" cy="17" r="2"/>',
    }
    path = paths.get(name, paths["file"])
    return f'<span class="icon icon-{color}"><svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24">{path}</svg></span>'


def section_header(icon_name, title, color="blue"):
    """Returns a modern section header with an icon."""
    return f'<div class="section-header">{icon(icon_name, color, 24)}<span>{title}</span></div>'


def show_skeleton_cards(count=4):
    """Shows N skeleton KPI cards while loading."""
    cols = st.columns(count)
    for col in cols:
        with col:
            st.markdown('<div class="skeleton skeleton-card"></div>', unsafe_allow_html=True)


def show_skeleton_rows(count=6):
    """Shows skeleton table rows while loading."""
    for _ in range(count):
        st.markdown('<div class="skeleton skeleton-row"></div>', unsafe_allow_html=True)


# ==================================================
# CONFIG
# ==================================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = "Steelworks_Fleet_Compliance.xlsm"

COMPLIANCE_SHEETS_ATTEMPT = {
    "SWPE": ["SWPE", "swpe", "Swpe", "SWPE ", " SWPE"],
    "SWIN": ["SWIN", "swin", "Swin", "SWIN ", " SWIN"],
}

FILE_PATH = os.path.join(BASE_DIR, EXCEL_FILE)

RC_COLUMN_LETTER = "H"

DOC_CONFIG = {
    "Insurance": {"expiry": "Insurance End Date",  "days": "Insurance Days Left", "col": "Q"},
    "Fitness":   {"expiry": "Fitness Expiry Date", "days": "Fitness Days Left",   "col": "V"},
    "MV Tax":    {"expiry": "MV Tax Expiry Date",  "days": "MV Tax Days Left",    "col": "AA"},
    "Permit":    {"expiry": "Permit Expiry Date",  "days": "Permit Days Left",    "col": "AE"},
    "TP":        {"expiry": "TP Expiry Date",      "days": "TP Days Left",        "col": "AL"},
}

SHAREPOINT_SEARCH_BASE = (
    "https://steelworkspower-my.sharepoint.com/personal/"
    "utkarsh_kashyap_steelworks_in/_layouts/15/onedrive.aspx?q="
)

DUE_DATE_ADVANCE_DAYS = 10


# ==================================================
# SESSION STATE
# ==================================================
if "dialog_open" not in st.session_state:
    st.session_state.dialog_open = False
if "dialog_vehicle" not in st.session_state:
    st.session_state.dialog_vehicle = None
if "search_query" not in st.session_state:
    st.session_state.search_query = ""
if "selected_month" not in st.session_state:
    st.session_state.selected_month = None


# ==================================================
# AUTO-DETECT SHEET NAMES
# ==================================================
@st.cache_data(ttl=60)
def get_actual_sheets():
    if not os.path.exists(FILE_PATH):
        return []
    try:
        with open(FILE_PATH, "rb") as f:
            data = f.read()
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        names = wb.sheetnames
        wb.close()
        return names
    except Exception:
        return []


def get_last_updated():
    if not os.path.exists(FILE_PATH):
        return None
    mtime = os.path.getmtime(FILE_PATH)
    return datetime.datetime.fromtimestamp(mtime)


# ==================================================
# LOAD HELPERS
# ==================================================
@st.cache_data(ttl=60)
def load_sheet(sheet_name):
    if not os.path.exists(FILE_PATH):
        return None
    for attempt in range(3):
        try:
            with open(FILE_PATH, "rb") as f:
                data = f.read()
            df = pd.read_excel(io.BytesIO(data), sheet_name=sheet_name, header=0, engine="openpyxl")
            df.columns = df.columns.str.strip()
            return df
        except PermissionError:
            time.sleep(1)
        except Exception:
            return None
    return None


@st.cache_data(ttl=60)
def load_hyperlinks(sheet_name):
    if not os.path.exists(FILE_PATH):
        return {}
    try:
        with open(FILE_PATH, "rb") as f:
            data = f.read()
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        ws = wb[sheet_name]

        links = {}
        for row in ws.iter_rows():
            for cell in row:
                url = None
                if cell.hyperlink and cell.hyperlink.target:
                    url = cell.hyperlink.target
                elif isinstance(cell.value, str) and cell.value.upper().startswith("=HYPERLINK"):
                    m = re.match(r'=HYPERLINK\(\s*"([^"]+)"', cell.value, re.IGNORECASE)
                    if m:
                        url = m.group(1)
                elif isinstance(cell.value, str) and cell.value.strip().lower().startswith(("http://", "https://")):
                    url = cell.value.strip()
                if url:
                    links.setdefault(cell.row, {})[cell.column_letter] = url
        return links
    except Exception:
        return {}


def clean_days_left(series):
    def parse(v):
        if pd.isna(v): return None
        if isinstance(v, (int, float)): return int(v)
        s = str(v).strip().upper()
        if "EXPIRED" in s:       return -1
        if s == "LTT":           return 99999
        if s == "ACTIVE":        return 9999
        if s == "EXPIRING SOON": return 15
        m = re.search(r"-?\d+", s)
        return int(m.group()) if m else None
    return series.apply(parse)


def status_icon(days):
    if days is None or (isinstance(days, float) and pd.isna(days)):
        return "⚪ N/A"
    if not isinstance(days, (int, float)):
        s = str(days).strip().upper()
        if "EXPIRED" in s: return "🔴 EXPIRED"
        if "LTT" in s:     return "🟢 LTT"
        if "ACTIVE" in s:  return "🟢 ACTIVE"
        m = re.search(r"-?\d+", s)
        if not m: return "⚪ N/A"
        days = int(m.group())
    d = int(days)
    if d < 0:   return "🔴 EXPIRED"
    if d <= 30: return "🟡 EXPIRING"
    return "🟢 ACTIVE"


def get_doc_url(doc_label, col_letter, hyperlinks, excel_row, vehicle, veh_col, sheet_name=None):
    url = hyperlinks.get(excel_row, {}).get(col_letter)
    if url:
        return url
    try:
        with open(FILE_PATH, "rb") as f:
            _data = f.read()
        _wb = openpyxl.load_workbook(io.BytesIO(_data), data_only=False)

        sheets_to_check = [sheet_name] if sheet_name else _wb.sheetnames

        for sn in sheets_to_check:
            if sn not in _wb.sheetnames:
                continue
            _ws = _wb[sn]
            cell = _ws[f"{col_letter}{excel_row}"]
            if cell.hyperlink and cell.hyperlink.target:
                return cell.hyperlink.target
            if isinstance(cell.value, str) and cell.value.upper().startswith("=HYPERLINK"):
                m = re.match(r'=HYPERLINK\(\s*"([^"]+)"', cell.value, re.IGNORECASE)
                if m:
                    return m.group(1)
            if isinstance(cell.value, str) and cell.value.strip().lower().startswith(("http://", "https://")):
                return cell.value.strip()
    except Exception:
        pass
    try:
        veh_no = str(vehicle[veh_col]).strip().replace(" ", "%20")
        return f"{SHAREPOINT_SEARCH_BASE}{doc_label}%20{veh_no}"
    except Exception:
        return None


def build_monthly_summary(df, months_ahead=12, advance_days=DUE_DATE_ADVANCE_DAYS):
    today = pd.Timestamp.today().normalize()
    cutoff = today + pd.DateOffset(months=months_ahead)

    summary = {}

    for doc_name, cfg in DOC_CONFIG.items():
        col = cfg["expiry"]
        if col not in df.columns:
            continue

        for _, row in df.iterrows():
            val = row.get(col)
            if pd.isna(val):
                continue
            try:
                expiry = pd.to_datetime(val)
            except Exception:
                continue

            due = expiry - pd.Timedelta(days=advance_days)
            if due < today or due > cutoff:
                continue

            key = (due.year, due.month)
            if key not in summary:
                summary[key] = {
                    "year": due.year,
                    "month": due.month,
                    "month_name": due.strftime("%B %Y"),
                    "total": 0,
                    "docs": {},
                    "days_left_min": None,
                    "items": [],
                }

            summary[key]["total"] += 1
            summary[key]["docs"][doc_name] = summary[key]["docs"].get(doc_name, 0) + 1

            days_left = (due - today).days
            if summary[key]["days_left_min"] is None or days_left < summary[key]["days_left_min"]:
                summary[key]["days_left_min"] = days_left

            summary[key]["items"].append({
                "Vehicle No": row.get(df.columns[0], ""),
                "Vehicle Name": row.get(df.columns[1], ""),
                "Document": doc_name,
                "Expiry Date": expiry,
                "Due Date": due,
                "Days Until Due": days_left,
            })

    return [summary[k] for k in sorted(summary.keys())]


# ==================================================
# SKELETON LOADER (shown while data loads)
# ==================================================
loading_placeholder = st.empty()

with loading_placeholder.container():
    st.markdown(f'<div class="section-header">{icon("truck", "blue", 24)}<span>Loading dashboard...</span></div>', unsafe_allow_html=True)
    show_skeleton_cards(4)
    st.markdown("<div style='height:24px'></div>", unsafe_allow_html=True)
    st.markdown('<div class="skeleton skeleton-title"></div>', unsafe_allow_html=True)
    show_skeleton_rows(6)


# ==================================================
# SIDEBAR
# ==================================================
with st.sidebar:
    st.markdown(f'### {icon("filter", "blue", 20)} Filters', unsafe_allow_html=True)

    search_query = st.text_input("🔎 Search Anything",
                                 value=st.session_state.search_query,
                                 placeholder="Vehicle no., name, chassis, engine, policy...",
                                 key="search_input")
    st.session_state.search_query = search_query

    st.markdown("<div style='height:8px'></div>", unsafe_allow_html=True)

    entity = st.radio("Select Entity", ["SWPE", "SWIN"], horizontal=True, key="entity_filter")

    actual_sheets = get_actual_sheets()

    real_sheet = None
    for candidate in COMPLIANCE_SHEETS_ATTEMPT[entity]:
        for s in actual_sheets:
            if s.strip().upper() == candidate.strip().upper():
                real_sheet = s
                break
        if real_sheet:
            break

    if real_sheet is None:
        st.error(f"❌ Sheet '{entity}' not found in file.")
        st.stop()

    df_all     = load_sheet(real_sheet)
    hyperlinks = load_hyperlinks(real_sheet)

    if df_all is None:
        st.error("❌ Could not load data. Please check the Excel file.")
        st.stop()

    if search_query.strip():
        q = search_query.strip().lower()
        mask = df_all.apply(
            lambda row: q in " ".join(str(v).lower() for v in row.values if pd.notna(v)),
            axis=1,
        )
        df_all = df_all[mask]

    doc_type = st.selectbox("Document Type", list(DOC_CONFIG.keys()), key="doc_filter")
    cfg = DOC_CONFIG[doc_type]

    if cfg["expiry"] in df_all.columns:
        df_all["_Expiry"]   = pd.to_datetime(df_all[cfg["expiry"]], errors="coerce")
        df_all["_DaysLeft"] = (df_all["_Expiry"] - pd.Timestamp.today()).dt.days
    else:
        df_all["_Expiry"]   = pd.NaT
        df_all["_DaysLeft"] = None

    if df_all["_DaysLeft"].isna().all() and cfg["days"] in df_all.columns:
        df_all["_DaysLeft"] = clean_days_left(df_all[cfg["days"]])

    df_all["_DaysLeft"] = pd.to_numeric(df_all["_DaysLeft"], errors="coerce")
    df_all["_Year"]  = df_all["_Expiry"].dt.year
    df_all["_Month"] = df_all["_Expiry"].dt.month

    years = ["All"] + sorted(df_all["_Year"].dropna().unique().astype(int).tolist())
    year  = st.selectbox("Year", years, key="year_filter")

    month_names = {1:"January", 2:"February", 3:"March", 4:"April", 5:"May", 6:"June",
                   7:"July", 8:"August", 9:"September", 10:"October", 11:"November", 12:"December"}
    month_sel = st.selectbox("Month", ["All"] + list(month_names.values()), key="month_filter")
    month = 0 if month_sel == "All" else [k for k, v in month_names.items() if v == month_sel][0]

    df_filtered = df_all.copy()
    if year != "All":
        df_filtered = df_filtered[df_filtered["_Year"] == year]
    if month != 0:
        df_filtered = df_filtered[df_filtered["_Month"] == month]

    st.divider()

    if st.button("🔄  Reset", use_container_width=True):
        st.cache_data.clear()
        st.session_state.dialog_open    = False
        st.session_state.dialog_vehicle = None
        st.session_state.search_query   = ""
        st.session_state.selected_month = None
        st.rerun()

    _dt = get_last_updated()
    if _dt is not None:
        st.markdown(
            f'<div class="small-updated">{icon("clock", "slate", 12)} Last updated<br>'
            f'<strong>{_dt.strftime("%d %b %Y, %I:%M %p")}</strong></div>',
            unsafe_allow_html=True,
        )


# Clear the skeleton loader once data is ready
loading_placeholder.empty()


# ==================================================
# HERO HEADER
# ==================================================
st.markdown(f"""
<div style="
    background: linear-gradient(135deg, #1e293b 0%, #0f172a 100%);
    border-radius: 16px;
    padding: 26px 30px;
    color: white;
    margin-bottom: 22px;
    box-shadow: 0 10px 30px -10px rgba(15,23,42,0.4);
    position: relative;
    overflow: hidden;
    animation: fadeInUp 0.5s cubic-bezier(0.4, 0, 0.2, 1);
">
    <div style="
        position: absolute;
        top: -50%;
        right: -10%;
        width: 320px;
        height: 320px;
        background: radial-gradient(circle, rgba(14,165,233,0.25) 0%, transparent 70%);
        border-radius: 50%;
    "></div>
    <div style="
        font-size: 11px;
        font-weight: 700;
        letter-spacing: 2px;
        opacity: 0.75;
        text-transform: uppercase;
        margin-bottom: 8px;
        display: flex;
        align-items: center;
        gap: 8px;
    ">
        {icon("shield", "blue", 14)}
        <span>Fleet Compliance</span>
    </div>
    <div style="font-size: 32px; font-weight: 900; letter-spacing: -1px; position: relative; z-index: 1;">
        {entity} Dashboard
    </div>
</div>
""", unsafe_allow_html=True)


# ==================================================
# KPI CARDS
# ==================================================
c1, c2, c3, c4 = st.columns(4)
with c1:
    st.markdown(f'<div style="font-size:11px;font-weight:700;color:#64748b;text-transform:uppercase;letter-spacing:1px;margin-bottom:4px;">{icon("truck", "slate", 12)} Total Vehicles</div>', unsafe_allow_html=True)
    st.metric(" ", len(df_filtered), label_visibility="collapsed")
with c2:
    st.markdown(f'<div style="font-size:11px;font-weight:700;color:#64748b;text-transform:uppercase;letter-spacing:1px;margin-bottom:4px;">{icon("check-circle", "green", 12)} Active</div>', unsafe_allow_html=True)
    st.metric(" ", int((df_filtered["_DaysLeft"] > 30).sum()), label_visibility="collapsed")
with c3:
    st.markdown(f'<div style="font-size:11px;font-weight:700;color:#64748b;text-transform:uppercase;letter-spacing:1px;margin-bottom:4px;">{icon("alert", "yellow", 12)} Expiring Soon</div>', unsafe_allow_html=True)
    st.metric(" ", int(((df_filtered["_DaysLeft"] >= 0) & (df_filtered["_DaysLeft"] <= 30)).sum()), label_visibility="collapsed")
with c4:
    st.markdown(f'<div style="font-size:11px;font-weight:700;color:#64748b;text-transform:uppercase;letter-spacing:1px;margin-bottom:4px;">{icon("x-circle", "red", 12)} Expired</div>', unsafe_allow_html=True)
    st.metric(" ", int((df_filtered["_DaysLeft"] < 0).sum()), label_visibility="collapsed")

st.markdown("<div style='height:24px'></div>", unsafe_allow_html=True)


# ==================================================
# 📅 MONTHLY DUE DATE BREAKDOWN
# ==================================================
st.markdown(
    section_header("calendar", "Monthly Due Date Breakdown", "blue"),
    unsafe_allow_html=True,
)
st.caption(f"Due date = Expiry Date − {DUE_DATE_ADVANCE_DAYS} days · Click any month to see the vehicle list")

monthly = build_monthly_summary(df_all, months_ahead=12, advance_days=DUE_DATE_ADVANCE_DAYS)

if not monthly:
    st.info("No documents due in the next 12 months.")
else:
    cols_per_row = 3
    for i in range(0, len(monthly), cols_per_row):
        row_items = monthly[i:i + cols_per_row]
        row_cols = st.columns(cols_per_row)

        for col_widget, item in zip(row_cols, row_items):
            with col_widget:
                days = item["days_left_min"] if item["days_left_min"] is not None else 999

                if days < 0: icon_char = "🔴"
                elif days <= 30: icon_char = "🟡"
                else: icon_char = "🟢"

                pills_text = " · ".join(f"{count} {doc}" for doc, count in item["docs"].items())
                btn_label = f"{icon_char}  {item['month_name']}   ({item['total']})\n{pills_text}"

                key = f"month_{item['year']}_{item['month']}"
                if st.button(btn_label, key=key, use_container_width=True):
                    st.session_state.selected_month = (item["year"], item["month"])
                    st.rerun()

    if st.session_state.selected_month is not None:
        sel_year, sel_month = st.session_state.selected_month
        sel_data = next((m for m in monthly if m["year"] == sel_year and m["month"] == sel_month), None)

        if sel_data:
            st.markdown("<div style='height:16px'></div>", unsafe_allow_html=True)
            st.markdown(
                section_header("file-text", f"Vehicles due in {sel_data['month_name']}", "violet"),
                unsafe_allow_html=True,
            )

            colA, colB = st.columns([4, 1])
            with colB:
                if st.button("✖  Close", key="close_month", use_container_width=True):
                    st.session_state.selected_month = None
                    st.rerun()

            items_df = pd.DataFrame(sel_data["items"])
            if not items_df.empty:
                items_df["Expiry Date"] = pd.to_datetime(items_df["Expiry Date"]).dt.strftime("%d-%b-%Y")
                items_df["Due Date"]    = pd.to_datetime(items_df["Due Date"]).dt.strftime("%d-%b-%Y")

                items_df = items_df[["Vehicle No", "Vehicle Name", "Document",
                                     "Expiry Date", "Due Date", "Days Until Due"]]

                st.dataframe(
                    items_df,
                    use_container_width=True,
                    hide_index=True,
                    column_config={
                        "Days Until Due": st.column_config.NumberColumn("Days Until Due", format="%d"),
                    },
                )

st.markdown("<div style='height:24px'></div>", unsafe_allow_html=True)
st.divider()


# ==================================================
# MAIN TABLE
# ==================================================
st.markdown(
    section_header("file", f"{doc_type} Records", "blue"),
    unsafe_allow_html=True,
)

if len(df_filtered) == 0:
    st.warning("⚠️ No vehicles match the current search or filters.")
    st.stop()

VEH_COL  = df_filtered.columns[0]
NAME_COL = df_filtered.columns[1]

display_cols = [VEH_COL, NAME_COL, cfg["expiry"], "_DaysLeft"]
available    = [c for c in display_cols if c in df_filtered.columns]

display_df = df_filtered[available].copy()
display_df.columns = ["Vehicle No", "Vehicle Name", "Expiry Date", "Days Left"][:len(available)]
display_df["Days Left"] = pd.to_numeric(display_df["Days Left"], errors="coerce")
display_df["Status"] = display_df["Days Left"].apply(status_icon)

header_cols = st.columns([2.2, 3.5, 2.2, 1.5, 1.8, 1.2])
header_cols[0].markdown("**Vehicle No**")
header_cols[1].markdown("**Vehicle Name**")
header_cols[2].markdown("**Expiry Date**")
header_cols[3].markdown("**Days Left**")
header_cols[4].markdown("**Status**")
header_cols[5].markdown("**Action**")

st.markdown("<hr style='margin:4px 0; border-color:#e2e8f0;'>", unsafe_allow_html=True)

for i, row in display_df.iterrows():
    cols = st.columns([2.2, 3.5, 2.2, 1.5, 1.8, 1.2])

    cols[0].write(row["Vehicle No"])
    cols[1].write(row["Vehicle Name"])

    exp_val = row["Expiry Date"]
    if pd.notna(exp_val):
        try:
            cols[2].write(pd.to_datetime(exp_val).strftime("%d-%b-%Y"))
        except Exception:
            cols[2].write(str(exp_val))
    else:
        cols[2].write("—")

    dl = row["Days Left"]
    cols[3].write(int(dl) if pd.notna(dl) else "—")
    cols[4].write(row["Status"])

    with cols[5]:
        btn_key = f"view_{entity}_{i}_{row['Vehicle No']}"
        if st.button("🔍 View", key=btn_key, use_container_width=True):
            st.session_state.dialog_vehicle = row["Vehicle No"]
            st.session_state.dialog_open    = True
            st.rerun()

    st.markdown("<hr style='margin:2px 0; border-color:#f1f5f9;'>", unsafe_allow_html=True)


# ==================================================
# VEHICLE DETAILS DIALOG
# ==================================================
if st.session_state.dialog_open and st.session_state.dialog_vehicle is not None:
    veh_no_clicked = str(st.session_state.dialog_vehicle).strip().upper()

    match_rows = df_filtered[
        df_filtered[VEH_COL].astype(str).str.strip().str.upper() == veh_no_clicked
    ]

    if match_rows.empty:
        st.session_state.dialog_open    = False
        st.session_state.dialog_vehicle = None
        st.rerun()

    vehicle = match_rows.iloc[0]

    excel_row = None
    try:
        with open(FILE_PATH, "rb") as f:
            _data = f.read()
        _wb = openpyxl.load_workbook(io.BytesIO(_data), data_only=True)
        _ws = _wb[real_sheet]
        for r in range(2, _ws.max_row + 1):
            cell_val = _ws.cell(row=r, column=1).value
            if cell_val and str(cell_val).strip().upper() == veh_no_clicked:
                excel_row = r
                break
    except Exception:
        pass

    if excel_row is None:
        excel_row = 2

    @st.dialog(f"🚛  {vehicle.get(NAME_COL, '')}  •  {vehicle[VEH_COL]}", width="large")
    def show_vehicle_details():

        st.markdown(f"""
        <div style="
            background: linear-gradient(135deg, #0ea5e9 0%, #0284c7 100%);
            border-radius: 14px;
            padding: 22px 26px;
            margin-bottom: 24px;
            color: white;
            box-shadow: 0 10px 30px -10px rgba(14,165,233,0.5);
            position: relative;
            overflow: hidden;
        ">
            <div style="
                position: absolute;
                top: -40%;
                right: -5%;
                width: 200px;
                height: 200px;
                background: radial-gradient(circle, rgba(255,255,255,0.15) 0%, transparent 70%);
                border-radius: 50%;
            "></div>
            <div style="font-size: 11px; font-weight: 700; opacity: 0.85; text-transform: uppercase; letter-spacing: 1.5px; margin-bottom: 6px; display: flex; align-items: center; gap: 6px; position: relative; z-index: 1;">
                {icon("truck", "slate", 12)}
                <span>Vehicle</span>
            </div>
            <div style="font-size: 24px; font-weight: 800; letter-spacing: -0.5px; position: relative; z-index: 1;">{vehicle[VEH_COL]}</div>
            <div style="font-size: 14px; opacity: 0.95; margin-top: 4px; position: relative; z-index: 1;">{vehicle.get(NAME_COL, '')}</div>
        </div>
        """, unsafe_allow_html=True)

        left, right = st.columns([1.7, 1])

        with left:
            st.markdown("### 📝 Vehicle Details")
            info = [
                ("RTO Location", "RTO Location"),
                ("Model No.", "Model No."),
                ("Chassis No.", "Chassis No."),
                ("Engine No", "Engine No"),
                ("Model Year", "Model Year"),
                ("Insurance Company", "Insurance Company"),
                ("Policy No.", "Policy No."),
                ("Insured Value (IDV)", "Insured Value (IDV)"),
                ("Insurance Amount", "Insurance Amount"),
                ("Fully Compliant Vehicle", "Fully Compliant Vehicle"),
            ]
            c1, c2 = st.columns(2)
            for i, (label, col) in enumerate(info):
                if col in vehicle.index and pd.notna(vehicle[col]):
                    val = vehicle[col]
                    if isinstance(val, float):
                        if col in ("Insured Value (IDV)", "Insurance Amount"):
                            val = f"₹ {val:,.0f}"
                        else:
                            val = f"{val:g}"
                    with (c1 if i % 2 == 0 else c2):
                        st.text_input(label, str(val), disabled=True, key=f"fld_{i}_{veh_no_clicked}")

        with right:
            st.markdown("### 📄 Documents")

            rc_url = get_doc_url("RC", RC_COLUMN_LETTER, hyperlinks, excel_row, vehicle, VEH_COL, real_sheet)
            if rc_url:
                st.link_button("📋   View RC", rc_url, use_container_width=True)
            else:
                st.button("❌   RC — No file", disabled=True,
                          use_container_width=True, key=f"nodoc_rc_{veh_no_clicked}")

            for doc_name, doc_cfg in DOC_CONFIG.items():
                doc_url = get_doc_url(doc_name, doc_cfg["col"], hyperlinks, excel_row, vehicle, VEH_COL, real_sheet)
                if doc_url:
                    st.link_button(f"📄   View {doc_name}", doc_url, use_container_width=True)
                else:
                    st.button(f"❌   {doc_name} — No file",
                              disabled=True, use_container_width=True,
                              key=f"nodoc_{doc_name}_{veh_no_clicked}")

        st.markdown("""
        <div style="text-align:center; color:#94a3b8; font-size:11px; padding:16px 0 4px 0; letter-spacing:0.3px;">
            Click outside or press <b>ESC</b> to close
        </div>
        """, unsafe_allow_html=True)

    show_vehicle_details()


# ==================================================
# EXPIRING SOON
# ==================================================
st.markdown("<div style='height:20px'></div>", unsafe_allow_html=True)
st.divider()
st.markdown(
    section_header("alert", "Expiring in 30 Days", "yellow"),
    unsafe_allow_html=True,
)

soon = df_filtered[(df_filtered["_DaysLeft"] >= 0) & (df_filtered["_DaysLeft"] <= 30)]

if soon.empty:
    st.success("✅ No documents expiring in the next 30 days.")
else:
    soon_cols  = [VEH_COL, NAME_COL, "_DaysLeft"]
    soon_avail = [c for c in soon_cols if c in soon.columns]
    soon_display = soon[soon_avail].copy()
    soon_display.columns = ["Vehicle No", "Vehicle Name", "Days Left"][:len(soon_avail)]
    st.dataframe(soon_display, use_container_width=True, hide_index=True)


# ==================================================
# FOOTER
# ==================================================
st.markdown(f"""
<div style="text-align:center; color:#94a3b8; font-size:12px; padding:20px 0 8px 0; margin-top:20px; border-top:1px solid #f1f5f9;">
    {icon("shield", "slate", 12)} &nbsp;© {pd.Timestamp.now().year} Steelworks • Fleet Compliance System
</div>
""", unsafe_allow_html=True)